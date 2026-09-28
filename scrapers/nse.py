"""
scrapers/nse.py - NSE Broker Locator -> LEAD_COLUMNS -> CSV + Excel.

    python -m scrapers.nse             # uses saved data/nse_debug.html (no re-download)
    python -m scrapers.nse --refresh   # download the ~78 MB page once more
    python -m scrapers.nse --max-members 20   # enrich only 20 members first (test run)
    python -m scrapers.nse --skip-enrich      # locator only, no member contacts

Flow:
    [1] Load page (saved file first; one download only if missing / --refresh)
    [2] Inspect displayNextResult()/displayPrevResult() in the page JS:
          client-side (show/hide rows)  -> pagination ignored, rows in HTML = full dataset
          server-side (AJAX/fetch/submit) -> request details printed, NOT followed
    [3] Parse EVERY <tr> of table#resultTable (lxml: fast on 78 MB)
    [4] Map to LEAD_COLUMNS (only values NSE shows; nothing invented)
    [5] Deduplicate on normalized name + company + address (segments merged)
    [6] Export output/trader_leads.csv + output/trader_leads.xlsx
    [7] Contact enrichment (before export), NSE "Know Your Broker" only:
          member listing -> member ID from each detail link ->
          /MemDirWeb/brokerDetailPage_Beta?h_MemType=members&memID=<ID>
          -> page kept only if it names that member -> member contact lookup
          -> merged into NSE Trading Member rows by exact normalised company name.
          Authorised Person rows are NEVER given their broker's contacts.
          Detail pages are cached in data/nse_member_contacts.json (resumable).

This script never solves, skips or bypasses a CAPTCHA or bot protection.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import time
from pathlib import Path
from urllib.parse import urljoin, urlparse

import pandas as pd
import requests
from lxml import html as lxml_html
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- config
SOURCE = "NSE"
URL = "https://enit.nseindia.com/MemDirWeb/searchBrokers_Beta?step=searchBrokersList"

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "output"
DEBUG_HTML = DATA_DIR / "nse_debug.html"
CSV_PATH = OUTPUT_DIR / "trader_leads.csv"
XLSX_PATH = OUTPUT_DIR / "trader_leads.xlsx"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-IN,en;q=0.9",
}

LEAD_COLUMNS = [
    "lead_id", "name", "designation", "company", "category",
    "email", "phone", "website",
    "linkedin", "twitter", "instagram", "youtube",
    "city", "state", "country",
    "market_segment", "registration_number",
    "source", "source_url",
]

NEXT_FN, PREV_FN = "displayNextResult", "displayPrevResult"
JS_KEYWORDS = {"function", "if", "for", "while", "return", "switch", "catch", "typeof", "parseInt", "String"}
SERVER_CALL_RE = re.compile(
    r"\$\.(?:ajax|post|get|getJSON)\s*\(|\bfetch\s*\(|XMLHttpRequest|\.submit\s*\(|"
    r"(?:window\.|document\.)?location(?:\.href)?\s*=|\.load\s*\(|\baxios\b",
    re.I,
)
CLIENT_VIEW_RE = re.compile(
    r"\.style\.display|\.show\s*\(|\.hide\s*\(|display\s*[:=]|visibility|\.toggle\s*\(|classList", re.I
)
URL_IN_JS_RE = re.compile(r"['\"]([^'\"\s]*(?:\.jsp|\.do|\.action|MemDirWeb|step=)[^'\"\s]*)['\"]")
CAPTCHA_RE = re.compile(r"captcha", re.I)
AP_COMPANY_RE = re.compile(r"^\s*(?:AP|Authori[sz]ed\s+Person)\s*(?:of|-|:)\s*(.+?)\s*$", re.I)
CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # Excel rejects these


def banner(title: str) -> None:
    print(f"\n{'=' * 80}\n{title}\n{'=' * 80}")


def clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", CONTROL_CHARS_RE.sub(" ", text or "")).strip()


def cell_text(el) -> str:
    """Text of a cell with <br>/nested tags separated by spaces."""
    return clean(" ".join(t.strip() for t in el.itertext() if t and t.strip()))


def norm_key(text: str) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


# --------------------------------------------------------------------------- [1] load
def load_page(refresh: bool) -> bytes:
    banner("[1] LOAD PAGE")
    if DEBUG_HTML.exists() and not refresh:
        print(f"Using saved page: {DEBUG_HTML} ({DEBUG_HTML.stat().st_size / 1e6:.1f} MB)")
        return DEBUG_HTML.read_bytes()

    print("Downloading NSE page once (large, may take a minute)...")
    resp = requests.get(URL, headers=HEADERS, timeout=300)
    print(f"Status: {resp.status_code} | {len(resp.content) / 1e6:.1f} MB")
    if resp.status_code in (401, 403, 429):
        raise SystemExit("Blocked by NSE (bot protection / rate limit). Not evading it. Stopping.")
    resp.raise_for_status()
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    DEBUG_HTML.write_bytes(resp.content)
    print(f"Saved to {DEBUG_HTML}")
    return resp.content


# --------------------------------------------------------------------------- [2] pagination
def extract_js_function(code: str, name: str) -> str | None:
    """Full function body via brace matching (a non-greedy regex stops at the first '}')."""
    m = re.search(
        r"function\s+" + re.escape(name) + r"\s*\(|" + re.escape(name) + r"\s*=\s*function\s*\(", code
    )
    if not m:
        return None
    start = code.find("{", m.end())
    if start == -1:
        return None
    depth = 0
    for i in range(start, min(len(code), start + 50000)):
        if code[i] == "{":
            depth += 1
        elif code[i] == "}":
            depth -= 1
            if depth == 0:
                return code[m.start(): i + 1]
    return code[m.start(): m.start() + 10000]


def _same_site_scripts(text: str) -> dict[str, str]:
    """Fetch same-site external JS (small files) only if the function is not inline."""
    out = {}
    for src in dict.fromkeys(re.findall(r"<script[^>]+src=[\"']([^\"']+)[\"']", text, re.I)):
        url = urljoin(URL, src)
        if urlparse(url).netloc != urlparse(URL).netloc:
            continue
        try:
            out[url] = requests.get(url, headers=HEADERS, timeout=30).text
        except requests.RequestException as exc:
            print(f"  could not fetch {url}: {exc}")
    return out


def inspect_pagination(text: str) -> str:
    """Return 'client', 'server' or 'unknown'. Nothing is followed."""
    banner("[2] PAGINATION MECHANISM (inspect only)")
    sources = {"page HTML": text}
    if extract_js_function(text, NEXT_FN) is None:
        sources.update(_same_site_scripts(text))

    bodies: dict[str, tuple[str, str]] = {}
    for fn in (NEXT_FN, PREV_FN):
        for where, code in sources.items():
            body = extract_js_function(code, fn)
            if body:
                bodies[fn] = (where, body)
                break

    # one level of helper functions called by next/prev
    for fn, (_, body) in list(bodies.items()):
        helpers = sorted(set(re.findall(r"\b([A-Za-z_]\w*)\s*\(", body)) - JS_KEYWORDS - set(bodies))
        for helper in helpers[:10]:
            for where, code in sources.items():
                hb = extract_js_function(code, helper)
                if hb:
                    bodies[helper] = (where, hb)
                    break

    if NEXT_FN not in bodies:
        print(f"{NEXT_FN}() not found -> mechanism UNKNOWN. Using the rows present in the HTML.")
        return "unknown"

    for fn, (where, body) in bodies.items():
        print(f"\n--- {fn}() [{where}] ---\n{body[:3000]}")

    combined = "\n".join(b for _, b in bodies.values())
    server = sorted(set(m.group(0) for m in SERVER_CALL_RE.finditer(combined)))
    client = sorted(set(m.group(0) for m in CLIENT_VIEW_RE.finditer(combined)))
    urls = sorted(set(URL_IN_JS_RE.findall(combined)))
    print(f"\nServer-request markers : {server or 'none'}")
    print(f"Client-view markers    : {client or 'none'}")
    print(f"URLs referenced        : {urls or 'none'}")

    if server:
        print("\nVERDICT: SERVER-SIDE. Lines that send the request:")
        for line in combined.splitlines():
            if SERVER_CALL_RE.search(line) or URL_IN_JS_RE.search(line):
                print("   ", line.strip()[:200])
        print(f"CAPTCHA referenced in pagination JS: {'YES' if CAPTCHA_RE.search(combined) else 'no'}")
        print("Pagination NOT followed. Only the rows already in this response are exported.")
        return "server"
    if client:
        print("\nVERDICT: CLIENT-SIDE. Next/Previous only change which rows are visible.")
        print("Pagination ignored - every <tr> in the HTML is the complete dataset.")
        return "client"
    print("\nVERDICT: UNKNOWN (no request and no show/hide found). Using rows present in the HTML.")
    return "unknown"


# --------------------------------------------------------------------------- [3] parse
def _column_indexes(headers: list[str]) -> dict[str, int] | None:
    low = [h.lower() for h in headers]

    def find(pred):
        return next((i for i, h in enumerate(low) if pred(h)), None)

    idx = {
        "name": find(lambda h: "name" in h),
        "address_type": find(lambda h: "address" in h and "type" in h),
        "address": find(lambda h: "address" in h and "type" not in h),
        "segment": find(lambda h: "segment" in h),
    }
    return idx if all(v is not None for v in idx.values()) else None


def parse_table(raw: bytes) -> tuple[list[dict], str | None]:
    banner("[3] PARSE table#resultTable (all rows)")
    t0 = time.time()
    doc = lxml_html.fromstring(raw)
    tables = doc.xpath("//table[@id='resultTable']")
    if not tables:
        raise SystemExit(
            "table#resultTable not found in this response (CAPTCHA page / blocked / changed layout). "
            "Not bypassing. Stopping."
        )
    table = tables[0]
    total_el = doc.xpath("//*[@id='totalRecordsId']")
    total_text = cell_text(total_el[0]) if total_el else None

    rows = table.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")  # no nested-table rows
    header_row = next((r for r in rows if r.xpath("./th")), None)
    headers = [cell_text(th) for th in header_row.xpath("./th")] if header_row is not None else []
    print(f"Headers        : {headers}")

    idx = _column_indexes(headers)
    if idx is None:
        idx = {"name": 0, "address_type": 1, "address": 2, "segment": 3}
        print("WARNING: headers not matched - using confirmed column order Name | Address Type | Address | Segment")
    need = max(idx.values())

    records, short, empty = [], 0, 0
    for tr in rows:
        tds = tr.xpath("./td")
        if not tds:
            continue
        if len(tds) <= need:
            short += 1
            continue
        rec = {k: cell_text(tds[i]) for k, i in idx.items()}
        if not rec["name"]:
            empty += 1
            continue
        records.append(rec)

    del doc
    print(f"<tr> in table  : {len(rows):,}")
    print(f"Rows parsed    : {len(records):,}")
    print(f"Skipped        : {short:,} too few cells | {empty:,} empty name")
    print(f"Total on page  : {total_text or 'not shown'}")
    print(f"Parse time     : {time.time() - t0:.1f}s")
    return records, total_text


# --------------------------------------------------------------------------- [4] map
def merge_segments(*values: str) -> str:
    parts: list[str] = []
    for v in values:
        for p in re.split(r"[,;/|]", v or ""):
            p = p.strip()
            if p and p not in parts:
                parts.append(p)
    return ", ".join(parts)


def to_lead(rec: dict) -> tuple[dict, str]:
    """Map one NSE row. Returns (lead, dedupe_key). Nothing is invented."""
    m = AP_COMPANY_RE.match(rec["address_type"])
    is_ap = m is not None
    company = clean(m.group(1)) if m else ""

    key = "|".join(norm_key(x) for x in (rec["name"], company, rec["address"]))  # key before TM company fill
    if not is_ap:
        company = rec["name"]  # a Trading Member row IS the member company
    lead = dict.fromkeys(LEAD_COLUMNS, "")
    lead.update({
        "lead_id": "NSE-" + hashlib.sha1(f"{SOURCE}|{key}".encode()).hexdigest()[:12],
        "name": rec["name"],
        "designation": "Authorized Person" if is_ap else "",
        "company": company,
        # Column is "Member/Authorised Person Name": a row that is not "AP of ..." is a member row.
        "category": "NSE Authorized Person" if is_ap else "NSE Trading Member",
        "country": "India",
        "market_segment": merge_segments(rec["segment"]),
        "source": SOURCE,
        "source_url": URL,
    })
    return lead, key


# --------------------------------------------------------------------------- [5] dedupe
def build_leads(records: list[dict]) -> tuple[pd.DataFrame, pd.DataFrame, int]:
    banner("[4-5] MAP + DEDUPLICATE (name + company + address)")
    unique: dict[str, dict] = {}
    duplicates = 0
    for rec in records:
        lead, key = to_lead(rec)
        if key in unique:
            duplicates += 1
            kept = unique[key]
            kept["lead"]["market_segment"] = merge_segments(kept["lead"]["market_segment"], lead["market_segment"])
            kept["raw"]["segment"] = kept["lead"]["market_segment"]
            if rec["address_type"] not in kept["raw"]["address_type"]:
                kept["raw"]["address_type"] += f" | {rec['address_type']}"
            continue
        unique[key] = {
            "lead": lead,
            "raw": {
                "lead_id": lead["lead_id"],
                "name": rec["name"],
                "address_type": rec["address_type"],
                "address": rec["address"],
                "segment": lead["market_segment"],
                "company": lead["company"],
            },
        }

    leads = pd.DataFrame([u["lead"] for u in unique.values()], columns=LEAD_COLUMNS)
    raw = pd.DataFrame([u["raw"] for u in unique.values()])
    print(f"Input rows      : {len(records):,}")
    print(f"Duplicates merged: {duplicates:,}")
    print(f"Unique records  : {len(leads):,}")
    if len(leads):
        print("Category counts :", leads["category"].value_counts().to_dict())
        print(f"With company    : {(leads['company'] != '').sum():,}")
    return leads, raw, duplicates


# --------------------------------------------------------------------------- [7] contact enrichment
MEMBERS_URL = "https://enit.nseindia.com/MemDirWeb/searchMembers_Beta?step=searchTradeMembersList"
MEMBERS_HTML = DATA_DIR / "nse_members.html"
MEMBER_CACHE = DATA_DIR / "nse_member_contacts.json"
MEMBER_DELAY_SECONDS = 1.5
MEMBER_RETRIES = 3

NSE_OWN_DOMAINS = ("nseindia.com", "nse.co.in", "nseit.com", "nsccl.co.in", "nseclearing.in")
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
SOCIAL_DOMAINS = {
    "linkedin": ("linkedin.com",),
    "twitter": ("twitter.com", "x.com"),
    "instagram": ("instagram.com",),
    "youtube": ("youtube.com", "youtu.be"),
}
# Exact labels first; "contains" rules below catch variants such as "Investor Grievance Email".
LABELS = {
    "member_name": {"member name", "name of the member", "trading member name", "name of member", "name"},
    "member_code": {"member code", "trading member code", "tm code", "member id"},
    "sebi_reg": {"sebi registration number", "sebi registration no", "sebi reg no", "sebi reg number",
                 "sebi registration", "registration number", "registration no"},
    "email": {"email", "email id", "e-mail", "e-mail id", "e mail", "email address"},
    "phone": {"phone", "phone no", "phone number", "telephone", "telephone no", "telephone number",
              "tel", "contact no", "contact number", "mobile", "mobile no", "mobile number"},
    "website": {"website", "website address", "web site", "url", "web address"},
    "city": {"city"},
    "state": {"state"},
}
# Contacts of a named individual (officer/director/...) are not the member's general contact.
PERSONAL_ROLE_RE = re.compile(r"officer|director|partner|proprietor|\bceo\b|\bcfo\b|\bmd\b|designated|authori[sz]ed", re.I)
CAPTCHA_HINT_RE = re.compile(r"captcha", re.I)


def label_kind(label: str) -> str | None:
    norm = re.sub(r"\s+", " ", re.sub(r"[:.*]+", " ", label.lower())).strip()
    if not norm or len(norm) > 60:
        return None
    for kind, names in LABELS.items():
        if norm in names:
            return kind
    if PERSONAL_ROLE_RE.search(norm):
        return None
    if "email" in norm or "e-mail" in norm:
        return "email"
    if any(w in norm for w in ("phone", "telephone", "mobile", "contact no", "contact number")):
        return "phone"
    if "website" in norm:
        return "website"
    if "sebi" in norm and "reg" in norm:
        return "sebi_reg"
    return None


def company_key(name: str) -> str:
    """Normalised company name for exact matching (LIMITED=LTD, PRIVATE=PVT, punctuation ignored)."""
    words = re.sub(r"[^A-Z0-9 ]+", " ", (name or "").upper().replace("&", " AND ")).split()
    swap = {"LIMITED": "LTD", "PRIVATE": "PVT", "COMPANY": "CO", "CORPORATION": "CORP"}
    words = [swap.get(w, w) for w in words if w != "THE"]
    return " ".join(words)


def _is_nse_own(text: str) -> bool:
    return any(d in text.lower() for d in NSE_OWN_DOMAINS)


def _uniq(values) -> list[str]:
    return list(dict.fromkeys(v for v in values if v))


def clean_emails(values) -> str:
    found = []
    for v in values:
        for e in EMAIL_RE.findall(v or ""):
            e = e.strip(".").lower()
            if not _is_nse_own(e):
                found.append(e)
    return "; ".join(_uniq(found))


def clean_phones(values) -> str:
    found = []
    for v in values:
        for part in re.split(r"[,;/|]", v or ""):
            part = clean(part)
            if len(re.sub(r"\D", "", part)) >= 7:  # ignore fragments / placeholders
                found.append(part)
    return "; ".join(_uniq(found))


def clean_website(values) -> str:
    for v in values:
        v = clean(v)
        if re.search(r"[a-z0-9-]+\.[a-z]{2,}", v, re.I) and "@" not in v and not _is_nse_own(v):
            return v
    return ""


def label_value_pairs(root) -> list[tuple[str, str]]:
    """Label/value pairs from tables (th/td, td/td), <dl>, 'Label : value' and label-line/value-line text."""
    pairs: list[tuple[str, str]] = []
    for tr in root.xpath(".//tr"):
        cells = [cell_text(c) for c in tr.xpath("./th|./td")]
        for i in range(0, len(cells) - 1, 2):
            if label_kind(cells[i]):
                pairs.append((cells[i], cells[i + 1]))
    for dl in root.xpath(".//dl"):
        for dt, dd in zip(dl.xpath("./dt"), dl.xpath("./dd")):
            pairs.append((cell_text(dt), cell_text(dd)))
    lines = [clean(t) for t in root.itertext()]
    lines = [t for t in lines if t and t != ":"]
    for i, line in enumerate(lines):
        if ":" in line:
            left, _, right = line.partition(":")
            if label_kind(left) and clean(right):
                pairs.append((left, clean(right)))
                continue
        if label_kind(line) and i + 1 < len(lines) and not label_kind(lines[i + 1]):
            pairs.append((line, lines[i + 1].lstrip(": ").strip()))
    return pairs


def parse_member_detail(html_text: str) -> dict:
    """Contact fields from ONE member's official detail page. Nothing is guessed."""
    doc = lxml_html.fromstring(html_text)
    for bad in doc.xpath("//script|//style|//noscript|//header|//footer|//nav"):
        bad.drop_tree()  # NSE's own header/footer contacts must never leak in
    buckets: dict[str, list[str]] = {}
    for label, value in label_value_pairs(doc):
        kind = label_kind(label)
        if kind and value:
            buckets.setdefault(kind, []).append(value)
    for a in doc.xpath("//a[starts-with(translate(@href,'MAILTO','mailto'),'mailto:')]"):
        buckets.setdefault("email", []).append(a.get("href", "")[7:])

    out = {
        "email": clean_emails(buckets.get("email", [])),
        "phone": clean_phones(buckets.get("phone", [])),
        "website": clean_website(buckets.get("website", [])),
        "member_code": clean((buckets.get("member_code") or [""])[0]),
        "sebi_reg": clean((buckets.get("sebi_reg") or [""])[0]).upper(),
        "member_name": clean((buckets.get("member_name") or [""])[0]),
        "city": clean((buckets.get("city") or [""])[0]),
        "state": clean((buckets.get("state") or [""])[0]),
    }
    for field_name, domains in SOCIAL_DOMAINS.items():
        links = [a.get("href", "") for a in doc.xpath("//a[@href]")]
        hit = next((h for h in links if any(d in h.lower() for d in domains) and "nse" not in h.lower()), "")
        out[field_name] = hit
    return out


def _nse_get(session: requests.Session, url: str) -> requests.Response:
    last = None
    for attempt in range(1, MEMBER_RETRIES + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=60)
            if resp.status_code in (401, 403, 429):
                raise PermissionError(f"HTTP {resp.status_code} (blocked / rate-limited)")
            resp.raise_for_status()
            return resp
        except requests.RequestException as exc:
            last = exc
            if attempt < MEMBER_RETRIES:
                time.sleep(MEMBER_DELAY_SECONDS * 2 * attempt)
    raise requests.RequestException(f"{url} failed after {MEMBER_RETRIES} attempts: {last}")


DETAIL_PATH = "/MemDirWeb/brokerDetailPage_Beta?h_MemType=members&memID={}"
MEMID_RE = re.compile(r"memID\s*[=:]\s*['\"]?([A-Za-z0-9_-]+)", re.I)
JS_CALL_RE = re.compile(r"([A-Za-z_$][\w$.]*)\s*\(([^)]*)\)")
NON_ID_ARGS = {"members", "member", "this", "event", "true", "false", "null", "undefined", ""}


def member_id_from_js(js: str) -> str | None:
    """NSE member ID from a detail link's JavaScript (memID=..., or the call's ID argument)."""
    m = MEMID_RE.search(js or "")
    if m:
        return m.group(1)
    for call in JS_CALL_RE.finditer(js or ""):
        if call.group(1).lower() in ("void", "javascript"):  # javascript:void(0) is not an ID
            continue
        args = [a.strip().strip("'\"") for a in call.group(2).split(",")]
        args = [a for a in args if a.lower() not in NON_ID_ARGS]
        numeric = [a for a in args if re.fullmatch(r"\d+", a)]
        if numeric:
            return numeric[0]
        ids = [a for a in args if re.fullmatch(r"[A-Za-z0-9_-]{1,40}", a)]
        if ids:
            return ids[0]
    return None


def detail_url_for(member_id: str) -> str:
    return urljoin(MEMBERS_URL, DETAIL_PATH.format(member_id))


def parse_member_listing(raw: bytes) -> list[dict]:
    """Know Your Broker listing: one dict per member, with its official detail-page URL."""
    doc = lxml_html.fromstring(raw)
    for table in doc.xpath("//table"):
        rows = table.xpath("./tr | ./thead/tr | ./tbody/tr")
        if len(rows) < 2:
            continue
        header_row = next((r for r in rows if r.xpath("./th")), rows[0])
        headers = [cell_text(c) for c in header_row.xpath("./th|./td")]
        kinds = [label_kind(h) or "" for h in headers]
        if "member_name" not in kinds or not ({"member_code", "sebi_reg"} & set(kinds)):
            continue
        members = []
        for tr in rows:
            if tr is header_row:
                continue
            tds = tr.xpath("./td")
            if len(tds) < len(headers):
                continue
            rec: dict = {}
            for kind, td in zip(kinds, tds):
                if kind and not rec.get(kind):
                    rec[kind] = cell_text(td)
            scripts = [tr.get("onclick") or ""]
            for a in tr.xpath(".//a"):
                href, onclick = (a.get("href") or "").strip(), (a.get("onclick") or "").strip()
                if href.lower().startswith("mailto:"):
                    rec["email"] = href[7:]
                elif href and not href.lower().startswith(("javascript", "#")):
                    rec.setdefault("detail_url", urljoin(MEMBERS_URL, href))
                scripts += [href if href.lower().startswith("javascript") else "", onclick]
            for inp in tr.xpath(".//input[@value]"):
                if "memid" in (inp.get("name", "") + inp.get("id", "")).lower():
                    scripts.append(f"memID={inp.get('value')}")
            js = " ".join(s for s in scripts if s)
            if js:
                rec["detail_js"] = js
                if not rec.get("detail_url"):
                    member_id = member_id_from_js(js)
                    if member_id:
                        rec["member_id"] = member_id
                        rec["detail_url"] = detail_url_for(member_id)
            if rec.get("member_name"):
                members.append(rec)
        return members
    return []


def _merge_contact(into: dict, other: dict) -> None:
    into["email"] = clean_emails([into.get("email", ""), other.get("email", "")])
    into["phone"] = clean_phones([into.get("phone", ""), other.get("phone", "")])
    for k, v in other.items():
        if k not in ("email", "phone") and not into.get(k):
            into[k] = v


def build_member_lookup(refresh: bool, max_members: int | None) -> tuple[dict[str, dict], list[str]]:
    """company_key -> member contact, from NSE's official member listing + detail pages."""
    banner("[7] CONTACT ENRICHMENT (NSE Know Your Broker)")
    notes: list[str] = []
    session = requests.Session()

    if MEMBERS_HTML.exists() and not refresh:
        raw = MEMBERS_HTML.read_bytes()
        print(f"Using saved member listing: {MEMBERS_HTML}")
    else:
        try:
            raw = _nse_get(session, MEMBERS_URL).content
        except (PermissionError, requests.RequestException) as exc:
            return {}, [f"Member listing not available: {exc}. Not bypassing."]
        MEMBERS_HTML.write_bytes(raw)
        print(f"Downloaded member listing ({len(raw) / 1e6:.1f} MB) -> {MEMBERS_HTML}")

    members = parse_member_listing(raw)
    if not members:
        hint = " (page asks for a CAPTCHA/search)" if CAPTCHA_HINT_RE.search(raw.decode("utf-8", "replace")) else ""
        return {}, [f"No member table in the Know Your Broker response{hint}. Not bypassing."]
    with_url = [m for m in members if m.get("detail_url")]
    print(f"Members listed      : {len(members):,}")
    print(f"With detail page URL: {len(with_url):,}")
    for m in with_url[:3]:
        print(f"  e.g. {m['member_name'][:40]!r}: {m.get('detail_js', '')[:60]!r} -> {m['detail_url']}")
    if not with_url:
        example = next((m["detail_js"] for m in members if m.get("detail_js")), "")
        notes.append(f"No member ID found in detail links (example: {example[:80]!r}). Listing columns only.")

    cache: dict[str, dict] = {}
    if MEMBER_CACHE.exists():
        try:
            cache = json.loads(MEMBER_CACHE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            notes.append("Contact cache was unreadable and has been rebuilt.")
    todo = [m for m in with_url if "name_on_page" not in cache.get(m["detail_url"], {})]
    if max_members is not None:
        todo = todo[:max_members]
    if todo:
        eta = len(todo) * (MEMBER_DELAY_SECONDS + 0.5) / 60
        print(f"Fetching {len(todo):,} detail pages (~{eta:.0f} min, already cached: {len(cache):,})")

    fetched_now: list[dict] = []
    headers = {**HEADERS, "Referer": MEMBERS_URL}
    for i, m in enumerate(todo, 1):
        time.sleep(MEMBER_DELAY_SECONDS)
        try:
            resp = session.get(m["detail_url"], headers=headers, timeout=60)
            if resp.status_code in (401, 403, 429):
                notes.append(f"Stopped after {i - 1} detail pages: HTTP {resp.status_code}. "
                             "Not bypassing; re-run later to resume.")
                break
            resp.raise_for_status()
        except requests.RequestException as exc:
            notes.append(f"Skipped {m['detail_url']}: {exc}")
            continue
        page = resp.text
        if CAPTCHA_HINT_RE.search(page) and "@" not in page:
            notes.append(f"Detail page asks for a CAPTCHA ({m['detail_url']}). Stopped; not bypassing.")
            break
        detail = parse_member_detail(page)
        page_text = company_key(" ".join(lxml_html.fromstring(page).itertext()))
        detail["name_on_page"] = company_key(m["member_name"]) in page_text  # right member's page?
        cache[m["detail_url"]] = detail
        fetched_now.append(detail)
        if i % 25 == 0 or i == len(todo):
            MEMBER_CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=0), encoding="utf-8")
            print(f"  detail pages: {i:,}/{len(todo):,}")
    if fetched_now:
        MEMBER_CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=0), encoding="utf-8")

    # ---- detail-page test statistics (this run)
    ok = [d for d in fetched_now if d.get("name_on_page")]
    has_e = sum(1 for d in ok if d["email"])
    has_p = sum(1 for d in ok if d["phone"])
    either = sum(1 for d in ok if d["email"] or d["phone"])
    banner("DETAIL PAGE RESULTS (this run)")
    print(f"Fetched detail pages       : {len(fetched_now):,}")
    print(f"Rejected (member not named): {len(fetched_now) - len(ok):,}")
    print(f"Contacts with email        : {has_e:,}")
    print(f"Contacts with phone        : {has_p:,}")
    print(f"Contacts with email OR phone: {either:,}")
    print(f"Contacts with neither      : {len(ok) - either:,}")

    # ---- member contacts: member code first, normalised name second
    by_code: dict[str, dict] = {}
    name_to_codes: dict[str, set[str]] = {}
    by_name_no_code: dict[str, dict] = {}
    rejected = 0
    for m in members:
        detail = cache.get(m.get("detail_url", ""), {})
        code = clean(m.get("member_code", ""))
        if detail and (not detail.get("name_on_page")
                       or (code and detail.get("member_code") and detail["member_code"] != code)):
            rejected += 1
            detail = {}  # page did not verifiably belong to this member
        contact = {
            "email": clean_emails([m.get("email", ""), detail.get("email", "")]),
            "phone": clean_phones([m.get("phone", ""), detail.get("phone", "")]),
            "website": clean_website([m.get("website", ""), detail.get("website", "")]),
            "sebi_reg": clean(m.get("sebi_reg") or detail.get("sebi_reg", "")).upper(),
            "member_code": code or detail.get("member_code", ""),
            "city": clean(m.get("city") or detail.get("city", "")),
            "state": clean(m.get("state") or detail.get("state", "")),
            **{k: detail.get(k, "") for k in SOCIAL_DOMAINS},
        }
        key = company_key(m["member_name"])
        if not key:
            continue
        code = contact["member_code"]
        if code:
            _merge_contact(by_code.setdefault(code, {}), contact)
            name_to_codes.setdefault(key, set()).add(code)
        else:
            _merge_contact(by_name_no_code.setdefault(key, {}), contact)

    lookup: dict[str, dict] = {}
    ambiguous = 0
    for key, codes in name_to_codes.items():
        if len(codes) == 1:
            lookup[key] = by_code[next(iter(codes))]
        else:
            ambiguous += 1  # one name, several member codes: assign none
    for key, contact in by_name_no_code.items():
        if key not in name_to_codes:
            lookup[key] = contact
    if rejected:
        notes.append(f"{rejected} detail pages ignored: member name/code on the page did not match the listing.")
    if ambiguous:
        notes.append(f"{ambiguous} member names map to several member codes; left unassigned.")
    usable = sum(1 for c in lookup.values() if c["email"] or c["phone"])
    print(f"\nUnique members for matching: {len(lookup):,} ({usable:,} with email or phone)")
    return lookup, notes


def apply_enrichment(leads: pd.DataFrame, lookup: dict[str, dict]) -> int:
    """Fill ONLY NSE Trading Member rows. Authorised Person rows are left untouched."""
    is_member = leads["category"] == "NSE Trading Member"
    matched = 0
    field_map = {"email": "email", "phone": "phone", "website": "website",
                 "registration_number": "sebi_reg", "city": "city", "state": "state",
                 "linkedin": "linkedin", "twitter": "twitter", "instagram": "instagram", "youtube": "youtube"}
    for idx in leads.index[is_member]:
        contact = lookup.get(company_key(leads.at[idx, "name"]))
        if not contact:
            continue
        matched += 1
        for col, src_key in field_map.items():
            if not str(leads.at[idx, col]).strip() and contact.get(src_key):
                leads.at[idx, col] = contact[src_key]
    return matched


def report(leads: pd.DataFrame, duplicates: int, matched: int, notes: list[str]) -> None:
    email = leads["email"].fillna("").str.strip().ne("")
    phone = leads["phone"].fillna("").str.strip().ne("")
    banner("CONTACT SUMMARY")
    print(f"Total NSE records                : {len(leads):,}")
    print(f"Trading Members                  : {(leads['category'] == 'NSE Trading Member').sum():,}")
    print(f"Authorized Persons               : {(leads['category'] == 'NSE Authorized Person').sum():,}")
    print(f"Trading Member rows matched      : {matched:,}")
    print(f"Records with email               : {email.sum():,}")
    print(f"Records with phone               : {phone.sum():,}")
    print(f"Records with email OR phone      : {(email | phone).sum():,}")
    print(f"Records with both email and phone: {(email & phone).sum():,}")
    print(f"Records with no contact          : {(~(email | phone)).sum():,}")
    print(f"Duplicates                       : {duplicates:,}")
    for n in notes:
        print(f"NOTE: {n}")


# --------------------------------------------------------------------------- [6] export
def _format_sheet(ws, df: pd.DataFrame) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        lengths = df[col].astype(str).str.len()
        typical = int(lengths.quantile(0.95)) if len(df) else 0
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(col), typical) + 2, 60)


def export(leads: pd.DataFrame, raw: pd.DataFrame) -> None:
    banner("[6] EXPORT")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    t0 = time.time()
    try:
        leads.to_csv(CSV_PATH, index=False, encoding="utf-8-sig")
        with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as writer:
            leads.to_excel(writer, sheet_name="leads", index=False)
            raw.to_excel(writer, sheet_name="nse_listing_raw", index=False)  # keeps address for the team
            _format_sheet(writer.sheets["leads"], leads)
            _format_sheet(writer.sheets["nse_listing_raw"], raw)
    except PermissionError as exc:
        raise SystemExit(f"{exc}\nClose the file in Excel and run again.") from exc
    print(f"Written in {time.time() - t0:.1f}s")


# --------------------------------------------------------------------------- main
def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass
    ap = argparse.ArgumentParser(description="NSE Broker Locator -> trader_leads.csv/.xlsx")
    ap.add_argument("--refresh", action="store_true", help="re-download instead of using data/nse_debug.html")
    ap.add_argument("--refresh-members", action="store_true", help="re-download the member listing")
    ap.add_argument("--skip-enrich", action="store_true", help="no contact enrichment")
    ap.add_argument("--max-members", type=int, default=None, help="fetch at most N new detail pages")
    args = ap.parse_args()

    raw_bytes = load_page(args.refresh)
    mode = inspect_pagination(raw_bytes.decode("utf-8", errors="replace"))
    records, total_text = parse_table(raw_bytes)
    del raw_bytes

    if not records:
        raise SystemExit("0 rows parsed from table#resultTable. Nothing exported.")
    leads, raw, duplicates = build_leads(records)

    matched, notes = 0, []
    if not args.skip_enrich:
        lookup, notes = build_member_lookup(args.refresh_members, args.max_members)
        matched = apply_enrichment(leads, lookup)
    export(leads, raw)

    banner("SUMMARY")
    print(f"Pagination      : {mode}")
    print(f"NSE rows parsed : {len(records):,}")
    print(f"NSE unique records: {len(leads):,}")
    print(f"CSV path        : {CSV_PATH}")
    print(f"Excel path      : {XLSX_PATH}")
    if total_text and re.sub(r"\D", "", total_text) and int(re.sub(r"\D", "", total_text)) != len(records):
        print(f"NOTE: page total says {total_text}, parsed {len(records):,} rows - check skipped counts in [3].")
    report(leads, duplicates, matched, notes)


if __name__ == "__main__":
    main()